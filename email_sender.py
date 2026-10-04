"""
Email Automation Tool
Sends personalized payment reminder emails to contacts from an Excel file.
"""

import smtplib
import ssl
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from openpyxl import load_workbook
import getpass


def read_contacts_from_excel(filename):
    """
    Read names and email addresses from an Excel file.

    Args:
        filename (str): Path to the Excel file (.xlsx)

    Returns:
        list: A list of dictionaries containing 'name' and 'email' for each contact
    """
    contacts = []

    # Load the workbook (Excel file)
    workbook = load_workbook(filename)

    # Get the active worksheet (the first sheet by default)
    worksheet = workbook.active

    # Iterate through each row in the worksheet
    # We start from row 2 (skipping header in row 1)
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        name = row[0]   # Column A (index 0)
        email = row[1]  # Column B (index 1)

        # Only add the contact if both name and email are present
        if name and email:
            contacts.append({
                'name': name,
                'email': email
            })

    return contacts


def create_email_message(recipient_name, recipient_email, sender_email):
    """
    Create a personalized email message.

    Args:
        recipient_name (str): The recipient's name
        recipient_email (str): The recipient's email address
        sender_email (str): The sender's email address

    Returns:
        MIMEMultipart: The constructed email message
    """
    # Create a multipart message (supports both plain text and HTML)
    message = MIMEMultipart("alternative")

    # Set email subject
    message["Subject"] = "Payment Reminder"

    # Set sender and recipient
    message["From"] = sender_email
    message["To"] = recipient_email

    # Create the email body with personalized greeting
    email_body = f"""Hi {recipient_name},

This is a friendly reminder that your payment is due. Please make the payment at your earliest convenience.

Thank you."""

    # Attach the plain text message
    # MIMEText converts our string into a format suitable for email
    text_part = MIMEText(email_body, "plain")
    message.attach(text_part)

    return message


def send_email(smtp_server, message, sender_email, app_password):
    """
    Send an email using Gmail's SMTP server.

    Args:
        smtp_server (smtplib.SMTP_SSL): The connected SMTP server
        message (MIMEMultipart): The email message to send
        sender_email (str): The sender's email address
        app_password (str): The Gmail App Password

    Returns:
        bool: True if sent successfully, False otherwise
    """
    try:
        # Send the email
        # as_string() converts the message object to a string format
        smtp_server.sendmail(sender_email, message["To"], message.as_string())
        return True
    except Exception as e:
        # Print the error and return False so we can continue with other emails
        print(f"  Error sending email: {e}")
        return False


def main():
    """
    Main function that orchestrates reading contacts and sending emails.
    """
    print("=" * 50)
    print("Email Automation Tool - Payment Reminder")
    print("=" * 50)
    print()

    # Configuration - Gmail SMTP settings
    SMTP_SERVER = "smtp.gmail.com"      # Gmail's SMTP server address
    SMTP_PORT = 465                      # Gmail's SSL port (secure connection)

    # Get sender's email credentials from user input
    # We use getpass for the password so it doesn't show on screen
    sender_email = input("Enter your Gmail address: ")
    app_password = getpass.getpass("Enter your Gmail App Password: ")
    print()

    # File containing contacts
    excel_file = "contacts.xlsx"

    print(f"Reading contacts from '{excel_file}'...")

    try:
        # Step 1: Read contacts from the Excel file
        contacts = read_contacts_from_excel(excel_file)
        print(f"Found {len(contacts)} contact(s).")
        print()
    except FileNotFoundError:
        print(f"Error: File '{excel_file}' not found.")
        print("Please make sure the file exists in the same folder as this script.")
        return
    except Exception as e:
        print(f"Error reading Excel file: {e}")
        return

    if not contacts:
        print("No contacts found in the file.")
        return

    # Create a secure SSL context (for encrypted connection)
    # This ensures our connection to Gmail is secure
    context = ssl.create_default_context()

    try:
        # Step 2: Connect to Gmail's SMTP server
        print("Connecting to Gmail SMTP server...")
        with smtplib.SMTP_SSL(SMTP_SERVER, SMTP_PORT, context=context) as server:
            # Login with our credentials
            print("Authenticating...")
            server.login(sender_email, app_password)
            print("Successfully connected!")
            print()

            # Step 3: Send email to each contact
            print("Sending emails...")
            print("-" * 50)

            success_count = 0
            fail_count = 0

            for contact in contacts:
                name = contact['name']
                email = contact['email']

                print(f"Sending to: {name} ({email})...")

                # Create the personalized email message
                message = create_email_message(name, email, sender_email)

                # Send the email
                if send_email(server, message, sender_email, app_password):
                    print(f"  ✓ Email sent successfully to {name}!")
                    success_count += 1
                else:
                    print(f"  ✗ Failed to send to {name}. Continuing...")
                    fail_count += 1

                print()

    except smtplib.SMTPAuthenticationError:
        print()
        print("Error: Authentication failed.")
        print("Please check your Gmail address and App Password.")
        print("Make sure you are using an App Password, not your regular password.")
        return
    except Exception as e:
        print(f"Error connecting to SMTP server: {e}")
        return

    # Step 4: Print summary
    print("-" * 50)
    print("Email sending complete!")
    print(f"  Successful: {success_count}")
    print(f"  Failed: {fail_count}")
    print("=" * 50)


# This block ensures main() only runs when the script is executed directly
# It won't run if this file is imported as a module
if __name__ == "__main__":
    main()
